"""Shared, PyQt5-free helpers used by more than one v2 module.

These are intentionally minimal: v2 modules take their working directories as
explicit function parameters rather than reading shared global state, so
nothing here is imported "implicitly" by other v2 modules. `setup_output_dirs`
exists purely as a CLI/demo convenience for wiring up a fresh output tree.
"""

import os

from openpyxl import load_workbook
import pandas as pd


def setup_output_dirs(base_output_dir: str) -> dict[str, str]:
    """Create (if needed) and return the standard v1-compatible output subfolders
    under `base_output_dir`: cell_classification_output, cell_track_output,
    quantitative_analysis_output. The data-export subfolder is intentionally
    omitted (that module is deferred).

    This is the single canonical replacement for two divergent v1
    implementations: `config.set_paths` (src/config.py:10-28, never actually
    called anywhere) and `gui_main.MainWindow.update_paths_and_create_folders`
    (src/gui_main.py:583-679), which duplicated the same logic inline and also
    set an undeclared `config.output_path` attribute. See v2/README.md's Bug
    Fix Changelog, entry #1.
    """
    paths = {
        "cell_classification_output": os.path.join(base_output_dir, "cell_classification_output"),
        "cell_track_output": os.path.join(base_output_dir, "cell_track_output"),
        "quantitative_analysis_output": os.path.join(base_output_dir, "quantitative_analysis_output"),
    }
    for path in paths.values():
        os.makedirs(path, exist_ok=True)
    return paths


def save_modified_rows(data: pd.DataFrame, file_path: str) -> pd.DataFrame:
    """Write rows flagged `modified == True` back into an existing .xlsx file in
    place (matching by 0-based DataFrame row position -> Excel row = position + 2),
    then clear the `modified` flag on the rows just saved. Returns `data` (mutated
    in place and returned for convenience).

    Ported verbatim from gui_comparison.py:239-265 (already pure/generic aside
    from its openpyxl I/O). Shared by the classification-correction and
    tracking-correction modules.
    """
    modified_rows = data[data["modified"] == True]  # noqa: E712 (matches v1's explicit bool comparison)

    if not modified_rows.empty:
        workbook = load_workbook(file_path)
        sheet = workbook.active

        col_index = {col_name: idx + 1 for idx, col_name in enumerate(data.columns)}

        for _, row in modified_rows.iterrows():
            excel_row = row.name + 2  # DataFrame row position + 2 (Excel is 1-based, row 1 is the header)
            for col_name, value in row.items():
                if col_name == "modified":
                    continue
                excel_col = col_index[col_name]
                sheet.cell(row=excel_row, column=excel_col, value=value)

        workbook.save(file_path)
        workbook.close()

        data.loc[modified_rows.index, "modified"] = False

    return data


def _self_test() -> None:
    """`save_modified_rows` maps columns BY POSITION, so widening the schema is
    the operation that can break it silently (Round 26 B4).

    It builds `{name: position}` from the DataFrame it is handed and writes each
    value at that position in the sheet. That is correct only while the sheet's
    column order equals the frame's -- and a correction round trip is exactly
    where the two meet: Extract writes the sheet, the correction UI re-reads it,
    edits rows, and writes them back into the same file. Round 26 appended
    `Edge Distance`, so this asserts the trip by NAME, over the real schema,
    with a value in the new last column that must not land anywhere else.
    """
    import tempfile

    import data_input

    columns = ["Cell Index"] + data_input.CELLS_INFO_COLUMNS
    rows = []
    for i in (1, 2, 3):
        row = {name: float(10 * i + j) for j, name in enumerate(data_input.CELLS_INFO_COLUMNS)}
        row["Cell Index"] = f"Cell1_{i}"
        row["Edge Distance"] = float(i)          # 1, 2, 3: the new LAST column
        rows.append(row)
    original = pd.DataFrame(rows, columns=columns)

    path = os.path.join(tempfile.mkdtemp(prefix="epiquant_common_"), "Cells_info.xlsx")
    original.to_excel(path, index=False, sheet_name="Cell")

    edited = pd.read_excel(path, sheet_name="Cell")
    assert list(edited.columns) == columns, list(edited.columns)
    edited["modified"] = False
    edited.loc[1, "Cell Area"] = 999.0
    edited.loc[1, "Edge Distance"] = 42.0        # the column most at risk
    edited.loc[1, "modified"] = True
    save_modified_rows(edited, path)

    back = pd.read_excel(path, sheet_name="Cell")
    assert list(back.columns) == columns, list(back.columns)
    assert back.loc[1, "Cell Area"] == 999.0, back.loc[1]
    assert back.loc[1, "Edge Distance"] == 42.0, back.loc[1]
    # The untouched rows and every other column of the edited row must be
    # unmoved -- an off-by-one would show up here, not in the two above.
    pd.testing.assert_frame_equal(back.drop(index=1), original.drop(index=1), check_dtype=False)
    unchanged = [c for c in columns if c not in ("Cell Area", "Edge Distance")]
    pd.testing.assert_series_equal(back.loc[1, unchanged], original.loc[1, unchanged], check_dtype=False)
    assert not edited["modified"].any(), "the flag must be cleared after saving"
    print("R26 B4: a correction round trip over the full %d-column schema puts both edits"
          % len(columns))
    print("  under their own names (including the appended `Edge Distance`), leaves every")
    print("  other cell byte-identical, and clears the modified flag.")
    print("\nAll common self-tests passed.")


if __name__ == "__main__":
    import argparse
    import tempfile

    parser = argparse.ArgumentParser(description="Standalone debug entry for v2.common")
    parser.add_argument("--base-output-dir", default=None, help="Defaults to a fresh temp dir if omitted")
    parser.add_argument("--self-test", action="store_true",
                        help="Run the schema round-trip check and exit.")
    args = parser.parse_args()

    if args.self_test:
        _self_test()
        raise SystemExit(0)

    base_dir = args.base_output_dir or tempfile.mkdtemp(prefix="epiquant_v2_")
    created = setup_output_dirs(base_dir)
    print(f"Base output dir: {base_dir}")
    for name, path in created.items():
        print(f"  {name}: {path} (exists={os.path.isdir(path)})")
