"""Module 4a: manual correction — cell classification (+ leading_edge computation).

Ported from `gui_comparison.MainWindow` (src/gui_comparison.py), the "Correct
Cell Classification Results" window. This file covers two related things:

1. The manual correction operations themselves (`update_clustering_data`,
   `update_cells_info_leading_edge`) — straightforward pure ports.
2. `compute_leading_edge` — a **new extraction**, not a mechanical hoist. In
   v1 this logic is interleaved across two ~150-line Qt-drawing methods,
   `update_image` (src/gui_comparison.py:391-556) and `update_image1`
   (558-739), which are near-duplicates but NOT equivalent:

   - `update_image` is the one actually used by the batch `process_images`
     job (src/gui_comparison.py:300) that computes and persists the
     `leading_edge` column into `Cells_info.xlsx` for every frame — this is
     the canonical, on-disk source of truth. It has BOTH branches: if any
     adjacent cell pairs are found for the frame, `leading_edge` = KDTree
     distance from each cell's center to the nearest separating point;
     otherwise a fixed sentinel (-1 for `cluster_label==1`, +1 for
     `cluster_label==0`).
   - `update_image1` is used only for interactive per-frame redraw while
     browsing; it recomputes+persists (via `update_cells_info`) ONLY in the
     "has matching points" branch — it has NO fallback branch at all, so for
     a frame with no adjacent pairs it silently leaves whatever value
     `update_image`/`process_images` already assigned untouched.

   `compute_leading_edge` below reproduces `update_image`'s (canonical,
   both-branch) semantics.

   Reproduced bug-for-bug (NOT one of the 4 agreed fixes, deliberately left
   as-is): in the "has matching points" branch, v1 assigns the *unsigned*
   KDTree distance regardless of cluster_label — there is no sign
   differentiation between the two compared clusters there, unlike the
   sentinel fallback branch which IS explicitly signed (-1/+1). Flagged here
   for your future consideration, not silently corrected.

N-cluster generalization: v1 (and this module until this change) hardcoded
`cluster_label == 1` as "the large-cell / -1-sentinel side" and `== 0` as
"the small-cell / +1-sentinel side" — a fixed convention tied to specific
integer values, not to which cluster has more member cells. (Confirmed by
running real data: for the classic AS/LE case, `Cluster_Label=1` is actually
the population MINORITY — "large cell" here meant physically large
individual cells, e.g. amnioserosa, which happen to be few in number, not
"the more populous cluster". Population-count and physical-cell-size are two
different things that happened to correlate in that one dataset.) Since
`pca_kmeans.cluster_cells_kmeans` now supports arbitrary N clusters with no
special "physically large cell population" concept at all, `compute_leading_edge`
takes an explicit `cluster_pair: tuple[int, int] | None` naming exactly two
`Cluster_Label` values to compare as the boundary's two sides:
`cluster_pair[0]` gets the -1 sentinel, `cluster_pair[1]` gets +1 — a purely
positional convention now, decided by the caller. For the classic 2-cluster
case (`cluster_pair=None`, exactly `{0, 1}` present), this auto-resolves to
`cluster_pair=(1, 0)` specifically to reproduce v1's exact numeric output
(since `cluster_cells_kmeans`'s new size-rank-descending numbering already
guarantees `Cluster_Label=0`=majority/`1`=minority for any 2-cluster run,
identical to v1's own numbering) — NOT a "whichever cluster has fewer cells"
runtime computation, which would be a coincidence specific to this one
dataset's biology, not a real generalizable rule.

`cells_info` throughout this module is expected to have `'Cell Index'` as a
plain column (not the DataFrame index) — this matches v1's actual runtime
shape, since `gui_comparison.MainWindow` loads `Cells_info.xlsx` via a plain
`pd.read_excel(path)` (no `index_col`), which turns the on-disk 'Cell Index'
index column back into a regular column on read. If you have a DataFrame
from `data_input.extract_cell_info` (which uses 'Cell Index' as the index),
call `.reset_index()` on it first.
"""

from typing import Optional

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.spatial import KDTree


def get_adjacent_cells(
    small_cells_contours: dict[int, np.ndarray],
    large_cells_contours: dict[int, np.ndarray],
    threshold: float = 3,
) -> list[tuple[int, int]]:
    """For every large-cell contour, KDTree-query every small-cell contour's
    points against it; if any point falls within `threshold`, the pair is
    "adjacent". Returns a list of `(large_cell_id, small_cell_id)` pairs.

    Verbatim port of `gui_comparison.MainWindow.get_adjacent_cells`
    (src/gui_comparison.py:361-376; already Qt-free). Parameter order matches
    v1 exactly: small-population contours first, large-population contours
    second — "small"/"large" here just label the two sides of whatever
    `cluster_pair` `compute_leading_edge` resolved, not any specific
    biological identity.
    """
    adjacent_pairs = []
    for large_cell_id, large_contour in large_cells_contours.items():
        large_tree = KDTree(large_contour)
        for small_cell_id, small_contour in small_cells_contours.items():
            distances, _ = large_tree.query(small_contour, k=1)
            if np.any(distances < threshold):
                adjacent_pairs.append((large_cell_id, small_cell_id))
    return adjacent_pairs


def get_matching_points_between_cells(
    large_contour: np.ndarray, small_contour: np.ndarray, threshold: float = 3,
) -> list[tuple[np.ndarray, np.ndarray]]:
    """For each point in `small_contour`, find its nearest point in
    `large_contour`; if within `threshold`, record the `(large_point,
    small_point)` pair. Verbatim port of
    `gui_comparison.MainWindow.get_matching_points_between_cells`
    (src/gui_comparison.py:378-389; already Qt-free).
    """
    matching_points = []
    large_tree = KDTree(large_contour)
    for small_point in small_contour:
        distance, index = large_tree.query(small_point, k=1)
        if distance < threshold:
            matching_points.append((large_contour[index], small_point))
    return matching_points


def compute_leading_edge(
    outlines: np.ndarray,
    frame_number: int,
    cells_info: pd.DataFrame,
    clustering_data: pd.DataFrame,
    threshold: float = 3,
    cluster_pair: Optional[tuple[int, int]] = None,
) -> pd.DataFrame:
    """Compute (or recompute) the `leading_edge` column for every cell in
    `frame_number`, given the frame's labeled segmentation mask `outlines`
    and the cell-type assignments in `clustering_data` (columns `Cell_Index`,
    `Cluster_Label`). Returns a copy of `cells_info` with `leading_edge`
    added/updated for that frame's rows.

    `cluster_pair` names exactly two `Cluster_Label` values to treat as the
    boundary's two sides: `cluster_pair[0]` gets the -1 sentinel (and feeds
    `large_cells_contours` internally), `cluster_pair[1]` gets +1. Cells
    belonging to any other `Cluster_Label` are skipped entirely (same as
    v1's behavior, generalized from "not in {0,1}" to "not in cluster_pair").

    If `cluster_pair` is `None`: auto-resolves only when `clustering_data`
    has exactly 2 distinct `Cluster_Label` values (else raises `ValueError`
    asking for an explicit pair). When those two values are exactly `{0, 1}`
    (the normal case for any 2-cluster `cluster_cells_kmeans` output),
    resolves to `(1, 0)` specifically to reproduce v1's exact numeric
    behavior — see the module docstring for why this is NOT the same as
    "whichever cluster has more cells". For any other pair of label values
    (e.g. after manual re-labeling via `update_clustering_data` leaves an
    unusual pair), resolves to the two values in ascending order — there is
    no legacy convention to preserve in that case, so it's just a
    deterministic default.

    Raises `ValueError` if an explicitly-given `cluster_pair` references a
    label not present in `clustering_data['Cluster_Label']`.

    See module docstring for the exact semantics ported (from
    `update_image`, the canonical batch-computation path) and the
    deliberately-unfixed unsigned-distance quirk.
    """
    unique_labels = set(clustering_data["Cluster_Label"].unique())

    if cluster_pair is None:
        if len(unique_labels) != 2:
            raise ValueError(
                f"clustering_data has {len(unique_labels)} distinct Cluster_Label values "
                f"({sorted(unique_labels)}); pass cluster_pair explicitly to select which two to compare."
            )
        cluster_pair = (1, 0) if unique_labels == {0, 1} else tuple(sorted(unique_labels))
    else:
        missing = [c for c in cluster_pair if c not in unique_labels]
        if missing:
            raise ValueError(f"cluster_pair value(s) {missing} not present in clustering_data['Cluster_Label'].")

    large_label, small_label = cluster_pair

    cell_type_dict = clustering_data.set_index("Cell_Index")["Cluster_Label"].to_dict()

    small_cells_contours: dict[int, np.ndarray] = {}
    large_cells_contours: dict[int, np.ndarray] = {}

    unique_ids = np.unique(outlines)
    for cell_id in unique_ids:
        if cell_id == 0:
            continue
        cell_index = f"Cell{frame_number}_{cell_id}"
        cluster_label = cell_type_dict.get(cell_index, None)
        if cluster_label == large_label:
            large_cells_contours[cell_id] = np.array(np.where(outlines == cell_id)).T
        elif cluster_label == small_label:
            small_cells_contours[cell_id] = np.array(np.where(outlines == cell_id)).T
        else:
            continue

    adjacent_pairs = get_adjacent_cells(small_cells_contours, large_cells_contours, threshold=threshold)

    all_matching_points = []
    for large_cell_id, small_cell_id in adjacent_pairs:
        large_contour = large_cells_contours[large_cell_id]
        small_contour = small_cells_contours[small_cell_id]
        matching_points = get_matching_points_between_cells(large_contour, small_contour, threshold=threshold)
        all_matching_points.extend([pt for large_point, pt in matching_points])

    cells_info = cells_info.copy()
    if "leading_edge" not in cells_info.columns:
        cells_info["leading_edge"] = None

    if all_matching_points:
        tree = KDTree(all_matching_points)
        for index, row in cells_info.iterrows():
            cell_index = row["Cell Index"]
            cluster_label = cell_type_dict.get(cell_index, None)
            if not cell_index.startswith(f"Cell{frame_number}_"):
                continue
            if cluster_label not in (large_label, small_label):
                continue
            cell_positions = np.where(outlines == int(cell_index.split("_")[-1]))
            center_y = int(np.mean(cell_positions[0]))
            center_x = int(np.mean(cell_positions[1]))
            cell_center = np.array([center_y, center_x])
            distance, _ = tree.query(cell_center)
            cells_info.at[index, "leading_edge"] = distance
    else:
        for index, row in cells_info.iterrows():
            cell_index = row["Cell Index"]
            cluster_label = cell_type_dict.get(cell_index, None)
            if not cell_index.startswith(f"Cell{frame_number}_"):
                continue
            if cluster_label not in (large_label, small_label):
                continue
            cells_info.at[index, "leading_edge"] = -1 if cluster_label == large_label else 1

    return cells_info


def find_boundary_cells(
    outlines: np.ndarray,
    frame_number: int,
    clustering_data: pd.DataFrame,
    threshold: float = 3,
    cluster_pair: Optional[tuple[int, int]] = None,
) -> set[str]:
    """`Cell_Index` strings for cells that directly touch (within
    `threshold` pixels) a cell of the *other* cluster in `cluster_pair` --
    the physical interface between the two populations, a thin
    one-cell-wide line (matches the manuscript's Fig. 2E leading-edge
    visualization).

    This is deliberately separate from `compute_leading_edge`'s stored
    `leading_edge` column: that column is a per-cell distance used for the
    mes/epi tracking split (`cell_tracking.split_by_leading_edge`), not
    itself an "is this cell on the boundary" flag, and a first cut at this
    function tried to derive one from it via a distance-percentile cutoff --
    which picks a broad, arbitrary-looking region rather than the actual
    touching cells. This version instead directly reuses `get_adjacent_cells`
    (the same pairwise-contour-proximity check `compute_leading_edge` already
    runs internally to find `all_matching_points`) and returns the cells
    that showed up in at least one adjacent pair. Pure visualization helper:
    doesn't touch `cells_info`, persists nothing, safe to call every overlay
    render.

    `cluster_pair` resolution mirrors `compute_leading_edge`'s exactly (only
    auto-resolves for a 2-cluster `clustering_data`, else pass explicitly;
    same `ValueError`s).
    """
    unique_labels = set(clustering_data["Cluster_Label"].unique())
    if cluster_pair is None:
        if len(unique_labels) != 2:
            raise ValueError(
                f"clustering_data has {len(unique_labels)} distinct Cluster_Label values "
                f"({sorted(unique_labels)}); pass cluster_pair explicitly to select which two to compare."
            )
        cluster_pair = (1, 0) if unique_labels == {0, 1} else tuple(sorted(unique_labels))
    else:
        missing = [c for c in cluster_pair if c not in unique_labels]
        if missing:
            raise ValueError(f"cluster_pair value(s) {missing} not present in clustering_data['Cluster_Label'].")

    large_label, small_label = cluster_pair
    cell_type_dict = clustering_data.set_index("Cell_Index")["Cluster_Label"].to_dict()

    small_cells_contours: dict[int, np.ndarray] = {}
    large_cells_contours: dict[int, np.ndarray] = {}
    for cell_id in np.unique(outlines):
        if cell_id == 0:
            continue
        cluster_label = cell_type_dict.get(f"Cell{frame_number}_{cell_id}", None)
        if cluster_label == large_label:
            large_cells_contours[cell_id] = np.array(np.where(outlines == cell_id)).T
        elif cluster_label == small_label:
            small_cells_contours[cell_id] = np.array(np.where(outlines == cell_id)).T

    adjacent_pairs = get_adjacent_cells(small_cells_contours, large_cells_contours, threshold=threshold)
    boundary_cell_ids = {cid for pair in adjacent_pairs for cid in pair}
    return {f"Cell{frame_number}_{cid}" for cid in boundary_cell_ids}


def update_clustering_data(clustering_data: pd.DataFrame, cell_index: str, classification_result: int) -> pd.DataFrame:
    """Manually reassign `cell_index` to cluster `classification_result` (any
    int — a valid `Cluster_Label` value, not restricted to 0/1): sets
    `Cluster_Label=classification_result`,
    `Cell_Type=f"cluster_{classification_result + 1}"`, and flags the row
    `modified=True` for `common.save_modified_rows` to pick up. Returns
    `clustering_data` (mutated in place and returned for convenience).

    N-cluster generalization: v1 only accepted `classification_result` 1
    ('mes Cell') or 0 ('epi Cell') — any other value silently did nothing
    except still setting `modified=True` (a latent no-op bug: the row got
    marked modified without its label actually changing). Since manual
    correction should be able to reassign a cell to any of the N clusters,
    not just a binary choice, that restriction is removed entirely here —
    every int is now a valid `classification_result`, so there's no
    remaining invalid-value case to (mis)handle. This does not validate that
    `classification_result` matches an existing cluster elsewhere in
    `clustering_data` — that's intentionally out of scope, not a bug; pass a
    sensible value.

    Ported from `gui_comparison.MainWindow.update_clustering_data`
    (src/gui_comparison.py:757-769). Requires a `'modified'` boolean column
    to already exist on `clustering_data` (v1 adds it at load time).
    """
    mask = clustering_data["Cell_Index"] == cell_index
    if mask.any():
        clustering_data.loc[mask, ["Cluster_Label", "Cell_Type"]] = [
            classification_result, f"cluster_{classification_result + 1}"
        ]
        clustering_data.loc[mask, "modified"] = True
    return clustering_data


def update_cells_info_leading_edge(cells_info: pd.DataFrame, cell_index: str, new_value: float) -> pd.DataFrame:
    """Set `leading_edge` for `cell_index` and flag the row `modified=True`.
    Adds the `leading_edge`/`modified` columns if missing (v1 does the same).
    Returns `cells_info` (mutated in place and returned for convenience).

    Ported from `gui_comparison.MainWindow.update_cells_info`
    (src/gui_comparison.py:782-798) — renamed here to disambiguate from
    `data_input`'s cell-info *extraction*.
    """
    if "leading_edge" not in cells_info.columns:
        cells_info["leading_edge"] = None

    mask = cells_info["Cell Index"] == cell_index
    if mask.any():
        cells_info.loc[mask, "leading_edge"] = new_value
        if "modified" not in cells_info.columns:
            cells_info["modified"] = False
        cells_info.loc[mask, "modified"] = True
    return cells_info


def save_leading_edge_column(cells_info: pd.DataFrame, file_path: str, sheet_name: str = "Cell") -> None:
    """Persist `cells_info['leading_edge']` back into an existing
    `Cells_info.xlsx` (written by `data_input.extract_cell_info`, which has
    no `leading_edge` column yet). Unlike `common.save_modified_rows` (which
    only ever updates existing cells, never adds a header), this writes the
    `'leading_edge'` header the first time the column doesn't exist yet on
    disk -- mirrors v1's one-time header-write in `process_images`
    (src/gui_comparison.py:310-330). `cells_info`'s row order/position must
    match the on-disk file's row order (true for the DataFrame
    `extract_cell_info` returns, as long as no rows were added/reordered/
    dropped in between -- `compute_leading_edge` never does).
    """
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.active

    header_row = [cell.value for cell in sheet[1]]
    if "leading_edge" in header_row:
        leading_edge_col = header_row.index("leading_edge") + 1
    else:
        leading_edge_col = sheet.max_column + 1
        sheet.cell(row=1, column=leading_edge_col, value="leading_edge")

    for index, row in cells_info.iterrows():
        excel_row = index + 2
        sheet.cell(row=excel_row, column=leading_edge_col, value=row["leading_edge"])

    workbook.save(file_path)
    workbook.close()


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Standalone debug entry for v2.manual_correction_classification")
    parser.add_argument("--npy-file", required=True, help="A single frame's .npy segmentation file")
    parser.add_argument("--frame-number", type=int, required=True)
    parser.add_argument("--cells-info", required=True, help="Cells_info.xlsx path")
    parser.add_argument("--clustering-data", required=True, help="cells_clustering_results.xlsx path")
    parser.add_argument(
        "--cluster-pair", type=int, nargs=2, default=None, metavar=("LABEL_A", "LABEL_B"),
        help="Two Cluster_Label values to compare (LABEL_A -> -1 sentinel, LABEL_B -> +1). "
             "Required if clustering-data has more than 2 distinct Cluster_Label values.",
    )
    args = parser.parse_args()

    dat = np.load(args.npy_file, allow_pickle=True).item()
    outlines = dat["outlines"]
    cells_info_df = pd.read_excel(args.cells_info)
    clustering_df = pd.read_excel(args.clustering_data)

    cluster_pair_arg = tuple(args.cluster_pair) if args.cluster_pair else None
    result = compute_leading_edge(
        outlines, args.frame_number, cells_info_df, clustering_df, cluster_pair=cluster_pair_arg
    )
    frame_mask = result["Cell Index"].str.startswith(f"Cell{args.frame_number}_")
    values = result.loc[frame_mask, "leading_edge"]
    print(f"cluster_pair used: {cluster_pair_arg if cluster_pair_arg else '(auto-resolved)'}")
    print(f"leading_edge for frame {args.frame_number}: n={len(values)}")
    print(f"  min={values.min()}, max={values.max()}")
    print(f"  count(+1)={(values == 1).sum()}, count(-1)={(values == -1).sum()}")
    print(f"  computed-distance count={((values != 1) & (values != -1)).sum()}")
